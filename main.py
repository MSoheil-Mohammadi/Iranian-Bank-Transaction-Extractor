"""
استخراج اطلاعات بانکی از فایل اکسل یا CSV تراکنش‌ها.

این ماژول فایل ورودی را می‌خواند، ستون «شرح تراکنش» را نرمال‌سازی
می‌کند و شماره کارت، شماره شبا و کد ملی را با اعتبارسنجی استخراج
می‌کند. سپس بانک صادرکننده را از روی BIN یا کد شبا تشخیص می‌دهد
و خروجی را به‌صورت اکسل (با چیدمان راست‌به‌چپ) یا CSV ذخیره می‌کند.

قابلیت‌ها:
    - خواندن فایل ورودی با فرمت xlsx/xlsm/xls یا csv
    - نوشتن فایل خروجی با فرمت xlsx یا csv
    - فیلتر کردن تراکنش‌های واریزی بر اساس ستون نوع تراکنش
    - اعتبارسنجی کارت با Luhn و شبا با Mod 97 و کد ملی با چک‌سام
    - تشخیص بانک از روی BIN کارت و کد بانک در شبا
    - استایل‌دهی ساده برای فایل‌های بزرگ (بیش از حد آستانه)

نحوه اجرا:
    python main.py input.xlsx Sheet1 "شرح تراکنش" output.xlsx
    python main.py input.csv "" "شرح تراکنش" output.csv --type-column "نوع"

پیش‌نیازها:
    - pandas
    - openpyxl
    - tqdm
    - پکیج‌های داخلی: utilities, constants
"""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from tqdm import tqdm

from constants import get_bank_by_bin, get_bank_by_iban_code
from utilities.iban import IBAN_PATTERN, clean_iban, iban_check
from utilities.luhn import CARD_PATTERN, clean_card, luhn_check
from utilities.national_id import (
    IRR_NATIONAL_CODE_PATTERN,
    NATIONAL_CODE_PATTERN,
    national_code_check,
)
from utilities.text import normalize_text

# ---------------------------------------------------------------------------
# ثابت‌ها
# ---------------------------------------------------------------------------

# آستانهٔ اعمال استایل کامل اکسل. برای فایل‌های بزرگ‌تر از این تعداد
# ردیف، استایل‌دهی جدول و تراز سلول‌ها اعمال نمی‌شود تا سرعت حفظ شود.
STYLE_THRESHOLD = 50_000


# ---------------------------------------------------------------------------
# توابع استخراج
# ---------------------------------------------------------------------------

def extract_valid_card(text: str) -> str | None:
    """استخراج اولین شماره کارت معتبر از متن.

    ابتدا با Regex کاندیدها پیدا می‌شوند، سپس هر کاندید با الگوریتم
    Luhn اعتبارسنجی می‌شود.

    Args:
        text: متن نرمال‌شده تراکنش.

    Returns:
        شماره کارت ۱۶ رقمی معتبر یا None.
    """
    for match in CARD_PATTERN.finditer(text):
        candidate = clean_card(match.group(0))
        if len(candidate) == 16 and luhn_check(candidate):
            return candidate
    return None


def extract_valid_iban(text: str) -> str | None:
    """استخراج اولین شماره شبا معتبر از متن.

    ابتدا با Regex کاندیدها پیدا می‌شوند، سپس هر کاندید با الگوریتم
    Mod 97 اعتبارسنجی می‌شود.

    Args:
        text: متن نرمال‌شده تراکنش.

    Returns:
        شماره شبا معتبر یا None.
    """
    for match in IBAN_PATTERN.finditer(text):
        candidate = clean_iban(match.group(0))
        if iban_check(candidate):
            return candidate
    return None


def extract_national_code(text: str) -> str | None:
    """استخراج اولین کد ملی معتبر از متن.

    ابتدا کدهایی که با IRR همراه هستند بررسی می‌شوند؛ در صورت نبود،
    همهٔ اعداد ۱۰ رقمی متن بررسی می‌شوند. هر کاندید با چک‌سام کد ملی
    اعتبارسنجی می‌شود.

    Args:
        text: متن نرمال‌شده (ارقام ASCII).

    Returns:
        کد ملی ۱۰ رقمی معتبر یا None.
    """
    # اولویت اول: کدهای همراه با IRR
    for match in IRR_NATIONAL_CODE_PATTERN.finditer(text):
        candidate = match.group(1) or match.group(2)
        if national_code_check(candidate):
            return candidate

    # اولویت دوم: هر عدد ۱۰ رقمی معتبر از نظر چک‌سام
    for match in NATIONAL_CODE_PATTERN.finditer(text):
        candidate = match.group(0)
        if national_code_check(candidate):
            return candidate

    return None


def extract_row(description) -> tuple:
    """استخراج همهٔ اطلاعات یک ردیف در یک فراخوانی.

    این تابع در یک پاس، متن را نرمال‌سازی می‌کند و هر سه شناسه
    (کارت، شبا، کد ملی) را استخراج و اعتبارسنجی می‌کند. همچنین
    بانک صادرکنندهٔ کارت و شبا را تشخیص می‌دهد.

    Args:
        description: مقدار خام ستون شرح تراکنش.

    Returns:
        تاپلی به شکل (شماره کارت، بانک کارت، شماره شبا، بانک شبا، کد ملی).
    """
    text = normalize_text(description)

    card = extract_valid_card(text)
    iban = extract_valid_iban(text)
    national_code = extract_national_code(text)

    card_bank = get_bank_by_bin(card) if card else None
    iban_bank = get_bank_by_iban_code(iban) if iban else None

    return card, card_bank, iban, iban_bank, national_code


# ---------------------------------------------------------------------------
# خواندن و نوشتن فایل
# ---------------------------------------------------------------------------

def read_input_file(input_file: str, sheet_name: str | None) -> pd.DataFrame:
    """خواندن فایل ورودی بر اساس پسوند آن.

    Args:
        input_file: مسیر فایل ورودی.
        sheet_name: نام شیت (فقط برای فایل‌های اکسل).

    Returns:
        دیتافریم خوانده‌شده با dtype=str.

    Raises:
        ValueError: اگر پسوند فایل پشتیبانی نشود.
        FileNotFoundError: اگر فایل وجود نداشته باشد.
    """
    path = Path(input_file)
    if not path.exists():
        raise FileNotFoundError(f"فایل ورودی یافت نشد: {input_file}")

    suffix = path.suffix.lower()

    if suffix == ".csv":
        return pd.read_csv(input_file, dtype=str, encoding="utf-8-sig")

    if suffix in (".xlsx", ".xlsm", ".xls"):
        return pd.read_excel(input_file, sheet_name=sheet_name, dtype=str)

    raise ValueError(f"فرمت ورودی پشتیبانی نمی‌شود: {suffix}")


def write_output_file(df: pd.DataFrame, output_file: str) -> None:
    """نوشتن فایل خروجی بر اساس پسوند آن.

    Args:
        df: دیتافریم خروجی.
        output_file: مسیر فایل خروجی.

    Raises:
        ValueError: اگر پسوند فایل پشتیبانی نشود.
    """
    suffix = Path(output_file).suffix.lower()

    if suffix == ".csv":
        df.to_csv(output_file, index=False, encoding="utf-8-sig")
        return

    if suffix == ".xlsx":
        _write_excel(df, output_file)
        return

    raise ValueError(f"فرمت خروجی پشتیبانی نمی‌شود: {suffix}")


def _write_excel(df: pd.DataFrame, output_file: str) -> None:
    """نوشتن فایل اکسل با چیدمان راست‌به‌چپ.

    برای فایل‌های کوچک‌تر از STYLE_THRESHOLD، جدول و تراز سلول‌ها
    اعمال می‌شود. برای فایل‌های بزرگ‌تر، فقط عرض ستون‌ها تنظیم
    می‌شود تا سرعت نوشتن حفظ شود.

    Args:
        df: دیتافریم خروجی.
        output_file: مسیر فایل اکسل خروجی.
    """
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Sheet1", index=False)
        ws = writer.sheets["Sheet1"]
        ws.sheet_view.rightToLeft = True

        # تنظیم عرض ستون‌ها (همیشه، چون سریع است)
        for col_idx, col_name in enumerate(df.columns, 1):
            width = max(len(str(col_name)) * 1.3, 12)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(width, 35)

        # اعمال استایل کامل فقط برای فایل‌های کوچک
        if len(df) <= STYLE_THRESHOLD:
            align = Alignment(horizontal="right", vertical="center")
            last_row = len(df) + 1
            last_col = get_column_letter(len(df.columns))
            table_ref = f"A1:{last_col}{last_row}"

            table = Table(displayName="Transactions", ref=table_ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            ws.add_table(table)

            for row in ws.iter_rows(
                min_row=2, max_row=last_row, max_col=len(df.columns)
            ):
                for cell in row:
                    cell.alignment = align

            ws.freeze_panes = "A2"
        else:
            print(
                f"⚠️ تعداد ردیف‌ها ({len(df):,}) از آستانه "
                f"({STYLE_THRESHOLD:,}) بیشتر است؛ استایل‌دهی ساده اعمال شد."
            )


# ---------------------------------------------------------------------------
# پردازش اصلی
# ---------------------------------------------------------------------------

def process_transactions(
    input_file: str,
    sheet_name: str,
    desc_column: str,
    output_file: str,
    type_column: str | None = None,
    deposit_keyword: str | None = None,
) -> None:
    """پردازش فایل ورودی و استخراج اطلاعات بانکی.

    Args:
        input_file: مسیر فایل ورودی (xlsx/xlsm/xls/csv).
        sheet_name: نام شیت (فقط برای اکسل).
        desc_column: نام ستون شرح تراکنش.
        output_file: مسیر فایل خروجی (xlsx/csv).
        type_column: نام ستون نوع تراکنش (اختیاری).
        deposit_keyword: کلیدواژهٔ واریز (اختیاری).

    Raises:
        ValueError: اگر ستون شرح تراکنش در فایل نباشد.
    """
    # خواندن فایل ورودی
    df = read_input_file(input_file, sheet_name)

    if desc_column not in df.columns:
        raise ValueError(f"ستون '{desc_column}' یافت نشد.")

    # فیلتر تراکنش‌های واریزی
    if type_column and deposit_keyword:
        if type_column in df.columns:
            mask = df[type_column].str.strip().str.contains(
                deposit_keyword.strip(), case=False, na=False
            )
            df = df[mask].copy()
            print(f"تعداد ردیف‌های واریزی: {len(df):,}")
        else:
            print(
                f"هشدار: ستون '{type_column}' یافت نشد. "
                f"تمام ردیف‌ها پردازش می‌شوند."
            )

    if df.empty:
        print("⚠️ هیچ ردیفی برای پردازش یافت نشد.")
        return

    # پردازش ردیف‌ها با نمایش نوار پیشرفت
    tqdm.pandas(desc="در حال پردازش ردیف‌ها")
    results = df[desc_column].progress_apply(extract_row)

    df["شماره کارت"] = [r[0] for r in results]
    df["بانک کارت"] = [r[1] for r in results]
    df["شماره شبا"] = [r[2] for r in results]
    df["بانک شبا"] = [r[3] for r in results]
    df["کد ملی"] = [r[4] for r in results]

    # نوشتن خروجی
    write_output_file(df, output_file)

    print("✅ استخراج با موفقیت انجام شد.")
    print(f"📁 خروجی: {output_file}")
    print(f"🔢 کارت‌های معتبر: {df['شماره کارت'].notna().sum():,}")
    print(f"🔢 شباهای معتبر: {df['شماره شبا'].notna().sum():,}")
    print(f"🔢 کدهای ملی: {df['کد ملی'].notna().sum():,}")


# ---------------------------------------------------------------------------
# اجرای برنامه از خط فرمان
# ---------------------------------------------------------------------------

def build_parser() -> argparse.ArgumentParser:
    """ساخت پارسر آرگومان‌های خط فرمان.

    Returns:
        پارسر آمادهٔ استفاده.
    """
    parser = argparse.ArgumentParser(
        description="استخراج اطلاعات بانکی از فایل اکسل یا CSV تراکنش‌ها"
    )
    parser.add_argument("input_file", help="فایل ورودی (xlsx/xlsm/xls/csv)")
    parser.add_argument(
        "sheet_name",
        help="نام شیت (برای فایل CSV خالی بگذار: \"\")",
    )
    parser.add_argument("desc_column", help="نام ستون شرح تراکنش")
    parser.add_argument("output_file", help="فایل خروجی (xlsx یا csv)")
    parser.add_argument(
        "--type-column",
        default=None,
        help="ستون نوع تراکنش (مثلاً 'نوع') – برای فیلتر واریز الزامی است",
    )
    parser.add_argument(
        "--deposit-keyword",
        default="واریز",
        help="کلیدواژهٔ واریز (پیش‌فرض: 'واریز')",
    )
    parser.add_argument(
        "--no-filter",
        action="store_true",
        help="نادیده گرفتن فیلتر واریز",
    )
    return parser


def main() -> None:
    """نقطهٔ ورود برنامه."""
    parser = build_parser()
    args = parser.parse_args()

    if args.no_filter:
        type_col = None
        deposit_kw = None
    else:
        type_col = args.type_column
        deposit_kw = args.deposit_keyword if type_col else None

    process_transactions(
        input_file=args.input_file,
        sheet_name=args.sheet_name,
        desc_column=args.desc_column,
        output_file=args.output_file,
        type_column=type_col,
        deposit_keyword=deposit_kw,
    )


if __name__ == "__main__":
    main()